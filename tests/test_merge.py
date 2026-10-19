import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from main import merge_tables


HEADERS = ['供应商编码', '供应商名称', '单子金额', '上货金额', '原因']


def write_book(path, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class MergeTablesTests(unittest.TestCase):
    def test_merges_and_keeps_rows_unique_to_each_input(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first, second, output = root / 'first.xlsx', root / 'second.xlsx', root / 'out.xlsx'
            write_book(first, [['A', '甲', 100, None, '首次'], ['B', '乙', 300, 30, None]])
            write_book(second, [['A', '甲', 200, 20, '补充'], ['C', '丙', 500, 50, None]])

            self.assertEqual(merge_tables(first, second, output), 3)
            workbook = load_workbook(output, data_only=True)
            rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
            workbook.close()
            self.assertEqual([row[0] for row in rows], ['A', 'B', 'C'])
            self.assertEqual(rows[0][2:5], ('100+200', 20, '首次+补充'))

    def test_rejects_duplicate_codes_in_second_input(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first, second, output = root / 'first.xlsx', root / 'second.xlsx', root / 'out.xlsx'
            write_book(first, [['A', '甲', 100, None, None]])
            write_book(second, [['A', '甲', 200, None, None], ['A', '重复', 300, None, None]])
            with self.assertRaisesRegex(ValueError, '重复供应商编码: A'):
                merge_tables(first, second, output)
            self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main()
